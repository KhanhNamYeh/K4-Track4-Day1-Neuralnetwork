"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (bằng code, không gõ tay).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu — step0_gap_vs_lnC, gap_val_minus_train, delta_val_f1_vs_base,
 beyond_noise — tự tính, không ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

LOSS_LABEL = {"ce": "CE", "mse": "MSE"}
OPT_LABEL = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
MAX_ROWS = 60  # số dòng có sẵn công thức trong mẫu (dòng 2..61)


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json. Trả về đường dẫn."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    payload = {k: result[k] for k in ("cfg", "history", "summary")}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=1, ensure_ascii=False)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json ở tầng trên cùng của results_dir, sắp theo exp_id."""
    out = []
    for p in sorted(Path(results_dir).glob("*.json")):
        with open(p, encoding="utf-8") as f:
            out.append(json.load(f))
    return out


def to_row(result: dict, eval_scores: dict | None = None, notes: str | None = None) -> dict:
    """Một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có) + figure_file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    c, s = result["cfg"], result["summary"]
    clip = c.get("clip_norm")
    row = dict(
        exp_id=c["exp_id"], group=c.get("group", "other"), description=c.get("description", ""),
        loss=LOSS_LABEL[c["loss"]], optimizer=OPT_LABEL[c["optimizer"]], lr=c["lr"],
        weight_decay=c["weight_decay"], batch=c["batch"], epochs=c["epochs"],
        hidden="-".join(str(h) for h in c["hidden"]), dropout=c["dropout"],
        clip_norm="none" if clip is None else clip, precision=c["precision"], init=c["init"], seed=c["seed"],
        step0_loss=s["step0_loss"], best_val_loss=s["best_val_loss"], best_epoch=s["best_epoch"],
        final_train_loss=s["final_train_loss"], final_val_loss=s["final_val_loss"],
        val_acc=s["val_acc"], val_macro_f1=s["val_macro_f1"], time_per_epoch_s=s["time_per_epoch_s"],
        peak_mem_MB=s["peak_mem_MB"], diverged=s["diverged"],
        eval_acc=None, eval_macro_f1=None, figure_file=f"figures/{c['exp_id']}.png",
        notes=notes if notes is not None else c.get("notes", ""),
    )
    if eval_scores:
        row["eval_acc"], row["eval_macro_f1"] = eval_scores["accuracy"], eval_scores["macro_f1"]
    return row


def _cell_value(v):
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               baseline_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu (từ dòng 2), ghi exp_id baseline vào sheet "Seeds",
    nhận xét theo nhóm vào cột H của sheet "Summary", rồi lưu thành out_path.

    Công thức của mẫu được giữ nguyên (không dùng data_only=True) và các cột công thức không bị ghi đè.
    Excel/LibreOffice tính lại công thức khi mở file.
    """
    import openpyxl
    if len(rows) > MAX_ROWS:
        raise ValueError(f"mẫu chỉ có sẵn công thức cho {MAX_ROWS} dòng, nhưng có {len(rows)} dòng")
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    header = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    formula_cols = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
    data_cols = [k for k in header if k not in formula_cols]
    for i in range(2, MAX_ROWS + 2):          # xoá dòng mẫu (chỉ các ô dữ liệu), giữ công thức/định dạng
        for k in data_cols:
            ws.cell(i, header[k]).value = None
    for i, row in enumerate(rows, start=2):
        for k in data_cols:
            ws.cell(i, header[k]).value = _cell_value(row.get(k))

    if baseline_ids is not None:
        seeds = wb["Seeds"]
        for i in range(2, 7):
            seeds.cell(i, 1).value = baseline_ids[i - 2] if i - 2 < len(baseline_ids) else None
    if summary_notes:
        sm = wb["Summary"]
        for r in range(2, 12):
            g = sm.cell(r, 1).value
            if g in summary_notes:
                sm.cell(r, 8).value = summary_notes[g]
    wb.save(out_path)
